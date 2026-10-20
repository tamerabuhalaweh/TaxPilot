"""XLSX renderer for the synthetic demo report (openpyxl, already pinned).

Produces a deterministic workbook with a mandatory synthetic-label row on every
sheet: ``Summary`` (fixture + policy + run metadata, the plain-language result
summary and counts), ``Findings`` (the accountant-facing review items) and
``Evidence`` (per-finding source-row pointers). No I/O beyond an in-memory
buffer.

REC-UI-02 presentation hardening: frozen header rows, keyboard-friendly
filters, readable column widths and wrapping, a result-summary block, and
formula-injection safety for every text cell (source text is data, never a
formula).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from app.demo.fixtures import SYNTHETIC_LABEL, DemoFixture
from app.demo.summary import Conclusion, build_conclusion, ordered_findings
from app.recon.gl import GLComparisonReport, GLFindingsReport

_BOLD = Font(bold=True)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="16437E")
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")

_ROLE_LABEL = {"sales": "Sales register", "gl": "General ledger"}


def _write(ws, row: int, column: int, value: object, *, font: Font | None = None):  # type: ignore[no-untyped-def]
    """Write a cell, forcing text to stay text (never an executable formula)."""
    cell = ws.cell(row=row, column=column)
    cell.value = value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        cell.data_type = "s"
    if font is not None:
        cell.font = font
    return cell


def _style_header(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _label_row(ws, columns: int) -> None:
    ws["A1"] = SYNTHETIC_LABEL
    ws["A1"].font = _BOLD
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)


def _currency_text(conclusion: Conclusion) -> str:
    parts: list[str] = []
    for result in conclusion.currency_results:
        if result.outcome == "DIFFERENCE":
            parts.append(
                f"{result.currency}: sales {result.sales_total}, GL {result.gl_total}, "
                f"difference {result.delta}"
            )
        else:
            parts.append(f"{result.currency}: coverage incomplete (no conclusion stated)")
    return "; ".join(parts) or "none"


def _summary_sheet(
    ws,
    fixture: DemoFixture,
    report: GLComparisonReport,
    findings: GLFindingsReport,
    engagement_period: tuple[str, str] | None,
) -> None:
    conclusion = build_conclusion(report, findings)
    rows: list[tuple[str, str]] = [
        ("Fixture", fixture.title),
        ("Fixture id", fixture.fixture_id),
        ("Jurisdiction", report.jurisdiction),
        ("Rules version", report.rules_version),
        (
            "Engagement period",
            f"{engagement_period[0]} to {engagement_period[1]}" if engagement_period else "not supplied",
        ),
        (
            "Fixture comparison period",
            f"{fixture.period_start.isoformat()} to {fixture.period_end.isoformat()}",
        ),
        ("Comparison basis", str(report.comparison_basis)),
        ("Sales measure", str(report.sales_measure)),
        ("Debit/credit convention", str(report.debit_credit_convention)),
        ("GL account scope", ", ".join(report.gl_account_scope)),
        ("Currency scope", ", ".join(report.currency_scope)),
        ("Tolerance", str(report.tolerance)),
        ("Blocked", str(report.blocked)),
        ("Inconclusive", str(report.inconclusive)),
        ("Inconclusive reasons", ", ".join(report.inconclusive_reasons)),
        ("Sales rows", str(report.sales_total)),
        ("GL rows", str(report.gl_total)),
        ("Review conclusion", conclusion.headline),
        ("Conclusion detail", conclusion.detail),
        ("Coverage complete", str(conclusion.coverage_complete)),
        ("Coverage currencies", ", ".join(conclusion.coverage_currencies) or "none"),
        ("Coverage explanations", "; ".join(conclusion.coverage_explanations) or "none"),
        ("Currency results", _currency_text(conclusion)),
        ("Caution", conclusion.caution),
        ("Report language", "English only (PDF and XLSX)"),
        ("Reconciliation fingerprint", report.fingerprint),
        ("Findings fingerprint", findings.fingerprint),
        ("Findings total", str(findings.summary.get("total", 0))),
        ("Findings review-required", str(findings.summary.get("review_required", 0))),
        ("Findings informational", str(findings.summary.get("informational", 0))),
        ("Synthetic label", SYNTHETIC_LABEL),
    ]
    _label_row(ws, 2)
    for offset, (key, value) in enumerate(rows, start=3):
        _write(ws, offset, 1, key, font=_BOLD).alignment = _TOP
        _write(ws, offset, 2, value).alignment = _WRAP
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 96
    ws.freeze_panes = "A3"


def _findings_sheet(ws, findings: GLFindingsReport) -> None:
    headers = [
        "#",
        "finding_id",
        "category",
        "severity",
        "outcome_kind",
        "review_required",
        "title",
        "what",
        "compared",
        "difference",
        "why",
        "guidance",
        "delta",
        "delta_currency",
        "rule",
        "rationale",
        "flags",
    ]
    _label_row(ws, len(headers))
    for col, name in enumerate(headers, start=1):
        _write(ws, 3, col, name)
    _style_header(ws, 3, len(headers))
    for index, item in enumerate(ordered_findings(findings.findings), start=1):
        row = 3 + index
        values: list[object] = [
            index,
            item.finding_id,
            item.category,
            item.severity,
            item.outcome_kind,
            str(item.review_required),
            item.title,
            item.what,
            item.compared,
            item.difference,
            item.why,
            item.guidance,
            item.delta or "",
            item.delta_currency or "",
            item.rule,
            "; ".join(item.rationale),
            ", ".join(item.flags),
        ]
        for col, value in enumerate(values, start=1):
            _write(ws, row, col, value).alignment = _WRAP
    widths = {
        1: 5,
        2: 22,
        3: 24,
        4: 10,
        5: 24,
        6: 10,
        7: 42,
        8: 44,
        9: 36,
        10: 46,
        11: 46,
        12: 50,
        13: 14,
        14: 10,
        15: 18,
        16: 40,
        17: 22,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + len(findings.findings)}"


def _evidence_sheet(ws, findings: GLFindingsReport) -> None:
    headers = ["finding_id", "role", "dataset_id", "source", "row_number", "raw_reference", "role_label"]
    _label_row(ws, len(headers))
    for col, name in enumerate(headers, start=1):
        _write(ws, 3, col, name)
    _style_header(ws, 3, len(headers))
    row = 4
    for item in ordered_findings(findings.findings):
        for ref in item.evidence:
            values: list[object] = [
                item.finding_id,
                ref.role,
                ref.dataset_id,
                ref.source,
                ref.row_number,
                ref.raw_reference,
                _ROLE_LABEL.get(ref.role, ref.role),
            ]
            for col, value in enumerate(values, start=1):
                _write(ws, row, col, value).alignment = _TOP
            row += 1
    for col, width in {1: 22, 2: 10, 3: 24, 4: 22, 5: 9, 6: 20, 7: 18}.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A4"
    last_row = max(3, row - 1)
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{last_row}"


def build_xlsx(
    fixture: DemoFixture,
    report: GLComparisonReport,
    findings: GLFindingsReport,
    *,
    engagement_period: tuple[str, str] | None = None,
) -> bytes:
    """Render the demo run into an in-memory XLSX workbook.

    ``engagement_period`` is display-only context: the calculation always uses
    the fixture's own comparison period, which the Summary sheet states
    explicitly so the two can never be conflated.
    """
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _summary_sheet(summary, fixture, report, findings, engagement_period)
    _findings_sheet(workbook.create_sheet("Findings"), findings)
    _evidence_sheet(workbook.create_sheet("Evidence"), findings)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["build_xlsx"]
