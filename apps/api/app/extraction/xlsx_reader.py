"""Safe read-only XLSX extraction (bounded ZIP pre-scan, no silent conversion).

Defense order: byte caps -> OOXML signature -> ZIP member scan (declared
totals) -> streaming actual-size verification -> read-only parse with
formula/external-link rejection -> explicit sheet/header selection ->
typed-cell policy (fail-closed by default, exact opt-in conversion).
Error messages carry coordinates and counts only, never cell values.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from app.extraction.contracts import (
    ExtractionError,
    ExtractionLimits,
    ExtractionResult,
    ExtractionStats,
    XlsxSettings,
)
from app.extraction.guards import check_bytes, check_doc_id, check_header_row
from app.normalization import ExtractedRow, ExtractedSheet, SourceDocument

_ZIP_MAGIC = b"PK\x03\x04"
_DRIVE_RE = re.compile(r"^[A-Za-z]:")
_MAX_LISTED_CELLS = 10
_CHUNK = 1024 * 1024


def _zip_scan(data: bytes, lim: ExtractionLimits) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionError("not a readable ZIP/OOXML container", code="MALFORMED_ZIP") from exc
    try:
        members = archive.infolist()
    except (zipfile.BadZipFile, NotImplementedError) as exc:
        raise ExtractionError("unreadable ZIP central directory", code="MALFORMED_ZIP") from exc
    if len(members) > lim.max_zip_members:
        raise ExtractionError(
            f"{len(members)} archive members exceeds {lim.max_zip_members}",
            code="ZIP_TOO_MANY_MEMBERS",
            detail={"members": len(members)},
        )
    names = [m.filename for m in members]
    if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise ExtractionError(
            "archive is not an OOXML workbook (missing content types/workbook)",
            code="NOT_XLSX",
        )
    declared_uncompressed = 0
    declared_compressed = 0
    for member in members:
        name = member.filename
        if name.startswith(("/", "\\")) or "\\" in name or _DRIVE_RE.match(name) or ".." in name.split("/"):
            raise ExtractionError("archive member with unsafe path", code="UNSAFE_MEMBER")
        lowered = name.lower()
        if lowered.endswith("vbaproject.bin"):
            raise ExtractionError("macros/VBA are never executed or extracted", code="MACROS_PRESENT")
        if lowered.startswith(("xl/embeddings/", "xl/activex/")) or "oleobject" in lowered:
            raise ExtractionError("embedded objects/OLE payloads are not extracted", code="EMBEDDED_OBJECTS")
        if lowered.startswith(("xl/externallinks", "xl/querytables", "xl/connections")):
            raise ExtractionError("external links/data connections are never followed", code="EXTERNAL_REFS")
        if member.flag_bits & 0x1:
            raise ExtractionError("password-protected entries are not supported", code="ENCRYPTED")
        if member.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
            raise ExtractionError("unsupported compression method", code="UNSUPPORTED_COMPRESSION")
        if member.is_dir():
            continue
        declared_uncompressed += member.file_size
        declared_compressed += member.compress_size
    if declared_uncompressed > lim.max_uncompressed_bytes:
        raise ExtractionError(
            f"archive declares {declared_uncompressed} uncompressed bytes "
            f"(limit {lim.max_uncompressed_bytes})",
            code="ZIP_TOO_LARGE",
            detail={"declared_uncompressed": declared_uncompressed},
        )
    if declared_uncompressed / max(declared_compressed, 1) > lim.max_compression_ratio:
        raise ExtractionError(
            f"archive compression ratio exceeds {lim.max_compression_ratio}:1",
            code="ZIP_BOMB_RATIO",
            detail={"declared_uncompressed": declared_uncompressed},
        )
    _verify_actual_sizes(archive, members, lim)
    return archive


def _verify_actual_sizes(
    archive: zipfile.ZipFile, members: list[zipfile.ZipInfo], lim: ExtractionLimits
) -> None:
    """Stream every member and bound real decompressed bytes (not just headers).

    Declared sizes are attacker-controlled; openpyxl would otherwise
    decompress unboundedly during parsing. Chunked reads keep memory flat.
    """
    total = 0
    try:
        for member in members:
            if member.is_dir():
                continue
            with archive.open(member) as stream:
                while True:
                    chunk = stream.read(_CHUNK)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > lim.max_uncompressed_bytes:
                        raise ExtractionError(
                            "actual decompressed bytes exceed the archive cap during " "verification read",
                            code="ZIP_BOMB",
                            detail={"verified_bytes": total},
                        )
    except RuntimeError as exc:  # e.g. password prompt mid-read
        raise ExtractionError(f"unreadable archive member: {exc}", code="MALFORMED_ZIP") from exc
    except zipfile.BadZipFile as exc:  # truncated/corrupt member (e.g. lied sizes)
        raise ExtractionError(
            "archive member failed integrity check during verification read",
            code="MALFORMED_ZIP",
        ) from exc


def _load_workbook(data: bytes) -> Any:
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ExtractionError(
            "malformed workbook: openpyxl cannot parse it read-only",
            code="MALFORMED_WORKBOOK",
        ) from exc


def _resolve_sheet_names(workbook: Any) -> list[str]:
    names = list(workbook.sheetnames)
    if not names:
        raise ExtractionError("workbook carries no worksheets", code="NO_WORKSHEETS")
    return names


def _select_sheet(names: list[str], settings: XlsxSettings) -> str:
    name = settings.sheet_name
    index = settings.sheet_index
    if name is not None and index is not None:
        indexed = names[index - 1] if 1 <= index <= len(names) else None
        if indexed != name:
            raise ExtractionError(
                "sheet_name and sheet_index disagree; refusing to guess",
                code="SHEET_CONFLICT",
            )
        assert indexed is not None
        return indexed
    if name is None and index is None:
        raise ExtractionError(
            "worksheet selection is required (sheet_name or 1-based sheet_index)",
            code="SHEET_REQUIRED",
        )
    if name is not None:
        if name not in names:
            raise ExtractionError(
                f"unknown worksheet; available: {', '.join(names[:50])}"
                + (f" (+{len(names) - 50} more)" if len(names) > 50 else ""),
                code="UNKNOWN_SHEET",
                detail={"worksheets": len(names)},
            )
        return name
    assert index is not None
    if not isinstance(index, int) or isinstance(index, bool) or not 1 <= index <= len(names):
        raise ExtractionError(
            "sheet_index out of range", code="UNKNOWN_SHEET", detail={"worksheets": len(names)}
        )
    return names[index - 1]


def _convert_cell(value: object, data_type: object) -> tuple[str | None, str | None]:
    """Return (text, refusal-kind). Refusals: 'formula', 'typed', 'error'."""
    if value is None:
        return None, None
    if data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        return None, "formula"
    if data_type == "e":
        return None, "error"
    if isinstance(value, str):
        return value, None
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "typed"
    if isinstance(value, int):
        return str(value), "typed"
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None, "error"
        return repr(value), "typed"
    if isinstance(value, datetime | date | time):
        return value.isoformat(), "typed"
    return None, "typed"


def _coord(column_index: int, row_number: int) -> str:
    return f"{get_column_letter(column_index)}{row_number}"


def extract_xlsx(
    data: bytes,
    *,
    doc_id: str,
    settings: XlsxSettings,
    limits: ExtractionLimits | None = None,
) -> ExtractionResult:
    """Extract one XLSX worksheet into ``ExtractedSheet`` + ``SourceDocument``.

    Fails closed on macros, external refs, formulas, error cells, and (unless
    explicitly opted in) typed numeric/date/boolean cells. See module docstring
    for the defense order.
    """
    lim = limits or ExtractionLimits()
    doc_id = check_doc_id(doc_id)
    if not isinstance(settings, XlsxSettings):
        raise TypeError(f"settings must be XlsxSettings, got {type(settings).__name__}")
    check_bytes(data, lim)
    if data[:4] != _ZIP_MAGIC:
        raise ExtractionError(
            "declared xlsx but content lacks the ZIP/OOXML signature "
            "(CSV text, legacy OLE .xls, or unknown format rejected)",
            code="NOT_XLSX",
        )
    header_row = check_header_row(settings.header_row)

    archive = _zip_scan(data, lim)
    with archive:
        workbook = _load_workbook(data)
        names = _resolve_sheet_names(workbook)
        selected = _select_sheet(names, settings)
        worksheet = workbook[selected]
        headers, numbered, skipped, scanned = _read_sheet(worksheet, selected, header_row, settings, lim)

    digest = hashlib.sha256(data).hexdigest()
    sheet = ExtractedSheet(headers=tuple(headers), rows=tuple(numbered))
    return ExtractionResult(
        doc=SourceDocument(doc_id=doc_id, content_hash=digest, worksheet=selected),
        sheet=sheet,
        stats=ExtractionStats(
            declared_format="xlsx",
            file_bytes=len(data),
            worksheet=selected,
            header_row=header_row,
            header_count=len(headers),
            data_rows=len(numbered),
            skipped_blank_rows=skipped,
            physical_rows_scanned=scanned,
        ),
    )


def _read_sheet(
    worksheet: Any,
    selected: str,
    header_row: int,
    settings: XlsxSettings,
    lim: ExtractionLimits,
) -> tuple[list[str], list[ExtractedRow], int, int]:
    skipped = 0
    scanned = 0
    headers: list[str] | None = None
    width = 0
    rows: list[ExtractedRow] = []

    for record in worksheet.iter_rows():
        if not record:
            continue
        row_number = record[0].row
        if row_number < header_row:
            continue
        texts: list[str] = []
        formulas: list[str] = []
        typed: list[tuple[str, str]] = []
        errors: list[str] = []
        for index, cell in enumerate(record, start=1):
            text, refusal = _convert_cell(cell.value, cell.data_type)
            coord = _coord(index, row_number)
            if refusal == "formula":
                formulas.append(coord)
            elif refusal == "error":
                errors.append(coord)
            elif refusal == "typed":
                typed.append((text or "", coord))
                texts.append(text or "")
            else:
                texts.append(text or "")
        # Fail the whole file at the first offending row (deterministic order).
        _raise_refusals(selected, settings, formulas, typed, errors)
        if row_number == header_row:
            while texts and texts[-1].strip() == "":
                texts.pop()  # trailing blanks never widen the schema
            headers = texts
            _check_headers(headers, lim)
            width = len(headers)
            continue
        if headers is None:  # sparse sheet without the header row: discard, fail below
            continue
        scanned += 1
        if scanned > lim.max_rows:
            raise ExtractionError(
                f"more than {lim.max_rows} data rows",
                code="TOO_MANY_ROWS",
                detail={"scanned": scanned},
            )
        while texts and texts[-1].strip() == "":
            texts.pop()
        if not texts:
            skipped += 1
            continue
        if len(texts) > width:
            raise ExtractionError(
                f"row {row_number} has {len(texts)} cells for {width} headers; "
                "kept as an error, never realigned by guessing",
                code="RAGGED_ROW",
                detail={"row_number": row_number},
            )
        while len(texts) < width:
            texts.append("")
        for index, cell in enumerate(texts):
            if len(cell) > lim.max_cell_chars:
                raise ExtractionError(
                    f"row {row_number} column {index} exceeds {lim.max_cell_chars} chars",
                    code="CELL_TOO_LARGE",
                    detail={"row_number": row_number, "column_index": index},
                )
        rows.append(ExtractedRow(row_number=row_number, cells=tuple(texts)))

    if headers is None:
        raise ExtractionError(
            f"header_row {header_row} not found in worksheet",
            code="HEADER_ROW_MISSING",
        )
    return headers, rows, skipped, scanned


def _check_headers(headers: list[str], lim: ExtractionLimits) -> None:
    if not headers or all(h.strip() == "" for h in headers):
        raise ExtractionError("header row carries no column names", code="EMPTY_HEADERS")
    if len(headers) > lim.max_columns:
        raise ExtractionError(
            f"{len(headers)} columns exceeds {lim.max_columns}",
            code="TOO_MANY_COLUMNS",
            detail={"columns": len(headers)},
        )
    for index, name in enumerate(headers):
        if name.strip() == "":
            raise ExtractionError("blank header cell", code="BLANK_HEADER", detail={"column_index": index})
        if len(name) > lim.max_cell_chars:
            raise ExtractionError(
                f"header at column {index} exceeds {lim.max_cell_chars} chars",
                code="CELL_TOO_LARGE",
                detail={"column_index": index},
            )


def _raise_refusals(
    selected: str,
    settings: XlsxSettings,
    formulas: list[str],
    typed: list[tuple[str, str]],
    errors: list[str],
) -> None:
    if formulas:
        raise ExtractionError(
            f"worksheet {selected!r} holds {len(formulas)} formula cells "
            f"({', '.join(formulas[:_MAX_LISTED_CELLS])}); cached values are "
            "never trusted, formulas are never evaluated",
            code="FORMULA_CELLS",
            detail={"cells": len(formulas)},
        )
    if errors:
        raise ExtractionError(
            f"worksheet {selected!r} holds {len(errors)} spreadsheet error values "
            f"({', '.join(errors[:_MAX_LISTED_CELLS])})",
            code="ERROR_VALUES",
            detail={"cells": len(errors)},
        )
    if typed and not settings.allow_typed_cells:
        raise ExtractionError(
            f"worksheet {selected!r} holds {len(typed)} typed numeric/date/boolean "
            f"cells ({', '.join(c for _, c in typed[:_MAX_LISTED_CELLS])}); refusing "
            "to silently convert display formatting into authoritative text "
            "(opt into allow_typed_cells for explicit exact conversion)",
            code="TYPED_CELLS",
            detail={"cells": len(typed)},
        )
