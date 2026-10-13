"""EP-04A XLSX extraction unit tests (synthetic workbooks built in-test)."""

import io
import re
import struct
import zipfile
from datetime import datetime
from typing import Any

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from app.extraction import (
    ExtractionError,
    ExtractionLimits,
    ExtractionResult,
    XlsxSettings,
    extract_bytes,
)


def make_xlsx(sheets: dict[str, list[list[Any]]]) -> bytes:
    """Build a synthetic .xlsx in memory (strings stay strings)."""
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet(name)
        if first:
            sheet.title = name
            first = False
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def extract(data: bytes, **kwargs: object) -> ExtractionResult:
    kwargs.setdefault("xlsx", XlsxSettings(sheet_name="Sales"))
    return extract_bytes(data, declared_format="xlsx", doc_id="SYN-XLSX-1", **kwargs)  # type: ignore[arg-type]


SALES_ROWS = [
    ["Invoice Number", "Invoice Date", "Net Amount", "Currency"],
    ["INV-0001", "05/01/2026", "1,234.500", "JOD"],
    ["000123", "06/01/2026", "-100.000", "JOD"],
]


def test_sales_sheet_happy_path() -> None:
    data = make_xlsx({"Sales": SALES_ROWS})
    result = extract(data)
    assert result.doc.worksheet == "Sales"
    assert result.sheet.headers == ("Invoice Number", "Invoice Date", "Net Amount", "Currency")
    assert [r.row_number for r in result.sheet.rows] == [2, 3]
    assert result.sheet.rows[1].cells[0] == "000123"
    assert result.stats.data_rows == 2 and result.stats.header_count == 4
    import hashlib

    assert result.doc.content_hash == hashlib.sha256(data).hexdigest()


def test_sheet_index_selection_and_header_row() -> None:
    data = make_xlsx({"Cover": [["Title"]], "Sales": SALES_ROWS})
    result = extract_bytes(data, declared_format="xlsx", doc_id="D", xlsx=XlsxSettings(sheet_index=2))
    assert result.doc.worksheet == "Sales" and result.stats.data_rows == 2
    titled = make_xlsx({"Sales": [["Report"], ["Generated"], *SALES_ROWS]})
    shifted = extract_bytes(
        data=titled,
        declared_format="xlsx",
        doc_id="D",
        xlsx=XlsxSettings(sheet_name="Sales", header_row=3),
    )
    assert shifted.sheet.headers[0] == "Invoice Number"
    assert shifted.sheet.rows[0].row_number == 4


def test_sheet_selection_errors() -> None:
    data = make_xlsx({"Sales": SALES_ROWS})
    with pytest.raises(ExtractionError) as exc:
        extract_bytes(data, declared_format="xlsx", doc_id="D", xlsx=XlsxSettings())
    assert exc.value.code == "SHEET_REQUIRED"
    with pytest.raises(ExtractionError) as exc:
        extract_bytes(data, declared_format="xlsx", doc_id="D", xlsx=XlsxSettings(sheet_name="Nope"))
    assert exc.value.code == "UNKNOWN_SHEET"
    with pytest.raises(ExtractionError) as exc:
        extract_bytes(data, declared_format="xlsx", doc_id="D", xlsx=XlsxSettings(sheet_index=9))
    assert exc.value.code == "UNKNOWN_SHEET"
    with pytest.raises(ExtractionError) as exc:
        extract_bytes(
            data,
            declared_format="xlsx",
            doc_id="D",
            xlsx=XlsxSettings(sheet_name="Other", sheet_index=1),
        )
    assert exc.value.code == "SHEET_CONFLICT"


def test_blank_rows_skipped_with_positions() -> None:
    rows = [SALES_ROWS[0], SALES_ROWS[1], ["", ""], SALES_ROWS[2]]
    result = extract(make_xlsx({"Sales": rows}))
    assert [r.row_number for r in result.sheet.rows] == [2, 4]
    assert result.stats.skipped_blank_rows == 1


def test_arabic_sheet_verbatim() -> None:
    rows = [["رقم الفاتورة", "صافي المبلغ"], ["000077", "300.00"]]
    data = make_xlsx({"المبيعات": rows})
    result = extract_bytes(data, declared_format="xlsx", doc_id="D", xlsx=XlsxSettings(sheet_name="المبيعات"))
    assert result.sheet.headers == ("رقم الفاتورة", "صافي المبلغ")
    assert result.sheet.rows[0].cells == ("000077", "300.00")


def test_typed_cells_fail_closed_by_default_without_value_leak() -> None:
    rows: list[list[Any]] = [["A", "B"], ["x", 987654321], ["y", datetime(2026, 1, 5)]]
    with pytest.raises(ExtractionError) as exc:
        extract(make_xlsx({"Sales": rows}))
    assert exc.value.code == "TYPED_CELLS"
    assert "987654321" not in str(exc.value)  # coordinates/counts only, never values


def test_typed_cells_opt_in_exact() -> None:
    rows: list[list[Any]] = [
        ["A", "B", "C", "D"],
        ["x", 42, 0.1, True],
        ["y", 7, datetime(2026, 1, 5, 9, 30), False],
    ]
    data = make_xlsx({"Sales": rows})
    result = extract(data, xlsx=XlsxSettings(sheet_name="Sales", allow_typed_cells=True))
    assert result.sheet.rows[0].cells == ("x", "42", "0.1", "TRUE")
    assert result.sheet.rows[1].cells == ("y", "7", "2026-01-05T09:30:00", "FALSE")


def test_formulas_refused_even_opted_in() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["A", "B"])
    sheet.append(["x", "=SUM(A1:A2)"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    settings = XlsxSettings(sheet_name="Sales", allow_typed_cells=True)
    with pytest.raises(ExtractionError) as exc:
        extract_bytes(buffer.getvalue(), declared_format="xlsx", doc_id="D", xlsx=settings)
    assert exc.value.code == "FORMULA_CELLS"
    assert "B2" in str(exc.value)


def test_error_values_refused() -> None:
    data = _patch_sheet_cell(
        make_xlsx({"Sales": [["A", "B"], ["x", "y"]]}),
        "B2",
        '<c r="B2" t="e"><v>#DIV/0!</v></c>',
    )
    with pytest.raises(ExtractionError) as exc:
        extract(data, xlsx=XlsxSettings(sheet_name="Sales", allow_typed_cells=True))
    assert exc.value.code == "ERROR_VALUES"


def _patch_sheet_cell(data: bytes, ref: str, replacement: str) -> bytes:
    """Rewrite one cell element inside xl/worksheets/sheet1.xml (synthetic)."""
    import re

    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        names = zin.namelist()
        target = next(n for n in names if n.endswith("worksheets/sheet1.xml"))
        xml = zin.read(target).decode("utf-8")
        patched, count = re.subn(rf'<c r="{ref}"[^>]*(/>|>.*?</c>)', replacement, xml, count=1)
        assert count == 1, f"cell {ref} not found for patching"
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                zout.writestr(name, patched if name == target else zin.read(name))
    return out.getvalue()


def _append_member(data: bytes, name: str, payload: bytes) -> bytes:
    """Return archive bytes with one extra member (synthetic attack fixture)."""
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        entries = [(item, zin.read(item.filename)) for item in zin.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item, content in entries:
            zout.writestr(item, content)
        zout.writestr(name, payload)
    return out.getvalue()


def test_external_links_macros_embeddings_rejected() -> None:
    base = make_xlsx({"Sales": SALES_ROWS})
    with pytest.raises(ExtractionError) as exc:
        extract(_append_member(base, "xl/externalLinks/ext1.xml", b"<x/>"))
    assert exc.value.code == "EXTERNAL_REFS"
    with pytest.raises(ExtractionError) as exc:
        extract(_append_member(base, "xl/vbaProject.bin", b"\x00" * 64))
    assert exc.value.code == "MACROS_PRESENT"
    with pytest.raises(ExtractionError) as exc:
        extract(_append_member(base, "xl/embeddings/oleObject1.bin", b"\x00" * 64))
    assert exc.value.code == "EMBEDDED_OBJECTS"
    with pytest.raises(ExtractionError) as exc:
        extract(_append_member(base, "../../evil.txt", b"evil"))
    assert exc.value.code == "UNSAFE_MEMBER"


def test_wrong_signatures_rejected() -> None:
    with pytest.raises(ExtractionError) as exc:
        extract(b"Invoice,Date\n1,2\n", xlsx=XlsxSettings(sheet_name="Sales"))
    assert exc.value.code == "NOT_XLSX"
    with pytest.raises(ExtractionError) as exc:
        extract(b"\xd0\xcf\x11\xe0" + b"\x00" * 100, xlsx=XlsxSettings(sheet_name="Sales"))
    assert exc.value.code == "NOT_XLSX"
    garbage = io.BytesIO()
    with zipfile.ZipFile(garbage, "w") as zf:
        zf.writestr("readme.txt", b"not a workbook")
    with pytest.raises(ExtractionError) as exc:
        extract(garbage.getvalue(), xlsx=XlsxSettings(sheet_name="Sales"))
    assert exc.value.code == "NOT_XLSX"


def test_zip_bomb_ratio_and_size_caps() -> None:
    rows = [["A"], *[["0" * 2000] for _ in range(30)]]
    data = make_xlsx({"Sales": rows})
    tiny = ExtractionLimits(max_uncompressed_bytes=1000, max_compression_ratio=2.0)
    with pytest.raises(ExtractionError) as exc:
        extract(data, limits=tiny)
    assert exc.value.code in ("ZIP_TOO_LARGE", "ZIP_BOMB_RATIO", "ZIP_BOMB")


def test_declared_sizes_lied_about_bomb() -> None:
    """Understate the sheet member's sizes: declared totals then pass, yet
    extraction still fails closed on integrity instead of succeeding with
    truncated garbage."""
    rows = [["A"], *[["0" * 100] for _ in range(3000)]]
    data = make_xlsx({"Sales": rows})
    with zipfile.ZipFile(io.BytesIO(data)) as check:
        (sheet_info,) = [i for i in check.infolist() if i.filename.endswith("worksheets/sheet1.xml")]
        assert sheet_info.file_size > 100000  # sanity: bulk really lives in the sheet member
    lied = _lie_about_member_sizes(data, "worksheets/sheet1.xml", 16)
    limits = ExtractionLimits(max_uncompressed_bytes=50000, max_compression_ratio=100000.0)
    with pytest.raises(ExtractionError) as exc:
        extract(lied, limits=limits)
    assert exc.value.code in ("ZIP_BOMB", "MALFORMED_ZIP", "MALFORMED_WORKBOOK")


def test_actual_size_verification_path_directly() -> None:
    """Honest-but-oversized content trips the streaming actual-size cap."""
    from app.extraction.xlsx_reader import _verify_actual_sizes

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("part.bin", b"0" * 100000)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        members = archive.infolist()
        with pytest.raises(ExtractionError) as exc:
            _verify_actual_sizes(archive, members, ExtractionLimits(max_uncompressed_bytes=1000))
    assert exc.value.code == "ZIP_BOMB"


def _lie_about_member_sizes(data: bytes, member_suffix: str, fake_size: int) -> bytes:
    """Understate one member's central-directory uncompressed size.

    Local headers stay intact so member offsets remain valid; only the
    declared total (attacker-controlled metadata our pre-scan reads) lies.
    """
    blob = bytearray(data)
    packed = struct.pack("<I", fake_size)
    text = bytes(blob)
    hits = 0
    start = 0
    while True:
        at = text.find(b"PK\x01\x02", start)
        if at < 0:
            break
        name_len = struct.unpack("<H", blob[at + 28 : at + 30])[0]
        name = bytes(blob[at + 46 : at + 46 + name_len]).decode("utf-8", errors="replace")
        if name.endswith(member_suffix):
            blob[at + 24 : at + 28] = packed  # central-dir uncompressed size
            hits += 1
        start = at + 4
    assert hits == 1, "target member not found exactly once for size patching"
    return bytes(blob)


def test_unsupported_compression_and_member_flood() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", b"<x/>", compress_type=zipfile.ZIP_BZIP2)
        zf.writestr("xl/workbook.xml", b"<x/>", compress_type=zipfile.ZIP_BZIP2)
    with pytest.raises(ExtractionError) as exc:
        extract(buffer.getvalue(), xlsx=XlsxSettings(sheet_name="S"))
    assert exc.value.code == "UNSUPPORTED_COMPRESSION"
    base = make_xlsx({"Sales": SALES_ROWS})
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as zin:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            for i in range(12):
                zout.writestr(f"xl/filler{i}.txt", b"pad")
    with pytest.raises(ExtractionError) as exc:
        extract(out.getvalue(), limits=ExtractionLimits(max_zip_members=10))
    assert exc.value.code == "ZIP_TOO_MANY_MEMBERS"


def test_shape_and_limit_errors() -> None:
    data = make_xlsx({"Sales": [["A", "B"], ["1", "2", "EXTRA"]]})
    with pytest.raises(ExtractionError) as exc:
        extract(data)
    assert exc.value.code == "RAGGED_ROW"
    assert exc.value.detail["row_number"] == 2
    short = make_xlsx({"Sales": [["A", "B", "C"], ["1", "2"]]})
    assert extract(short).sheet.rows[0].cells == ("1", "2", "")
    with pytest.raises(ExtractionError) as exc:
        extract(
            make_xlsx({"Sales": [["A", "B"], ["1", "2"], ["3", "4"]]}), limits=ExtractionLimits(max_rows=1)
        )
    assert exc.value.code == "TOO_MANY_ROWS"
    with pytest.raises(ExtractionError) as exc:
        extract(
            make_xlsx({"Sales": [["A", "B", "C"], ["1", "2", "3"]]}), limits=ExtractionLimits(max_columns=2)
        )
    assert exc.value.code == "TOO_MANY_COLUMNS"
    with pytest.raises(ExtractionError) as exc:
        extract(make_xlsx({"Sales": [["", "B"], ["1", "2"]]}))
    assert exc.value.code == "BLANK_HEADER"
    with pytest.raises(ExtractionError) as exc:
        extract(make_xlsx({"Sales": [["A", "B"]]}), xlsx=XlsxSettings(sheet_name="Sales", header_row=9))
    assert exc.value.code == "HEADER_ROW_MISSING"


def test_stream_and_file_caps() -> None:
    data = make_xlsx({"Sales": SALES_ROWS})
    with pytest.raises(ExtractionError) as exc:
        extract(data, limits=ExtractionLimits(max_file_bytes=10))
    assert exc.value.code == "FILE_TOO_LARGE"
    with pytest.raises(ExtractionError) as exc:
        extract(b"")
    assert exc.value.code == "EMPTY_FILE"
    from app.extraction import extract_stream

    via_stream = extract_stream(
        io.BytesIO(data), declared_format="xlsx", doc_id="S", xlsx=XlsxSettings(sheet_name="Sales")
    )
    assert via_stream.sheet == extract(data).sheet


# R1: worksheet dimension metadata is verified against actual content.
def _patch_sheet_xml(data: bytes, old: str, new: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        names = zin.namelist()
        target = next(n for n in names if n.endswith("worksheets/sheet1.xml"))
        xml = zin.read(target).decode("utf-8")
        patched, count = re.subn(old, new, xml, count=1)
        assert count == 1, "patch anchor not found exactly once"
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, patched if item.filename == target else zin.read(item.filename))
    return out.getvalue()


# openpyxl/et_xmlfile serialize empty elements with or without a space
# (`<dimension ref="A1:B2"/>` vs `<dimension ref="A1:B2" />`) depending on
# the installed XML backend, so test patching stays serialization-agnostic.
_DIMENSION_TAG_RE = r'<dimension\s+ref="[^"]*"\s*/?>'


def _patch_dimension(data: bytes, newref: str) -> bytes:
    return _patch_sheet_xml(data, _DIMENSION_TAG_RE, f'<dimension ref="{newref}"/>')


def test_forged_narrow_dimension_refused_not_silently_truncated() -> None:
    data = make_xlsx({"Sales": [["A", "B"], ["1", "2"], ["3", "4"]]})
    forged = _patch_dimension(data, "A1:B2")  # hides physical row 3 from dimension readers
    with pytest.raises(ExtractionError) as exc:
        extract(forged)
    assert exc.value.code == "DIMENSION_UNDERSTATED"


def test_forged_wide_dimension_refused_before_expansion() -> None:
    data = make_xlsx({"Sales": [["A", "B"], ["1", "2"]]})
    forged = _patch_dimension(data, "A1:XFD1000000")  # 16k cols x 1M rows phantom
    with pytest.raises(ExtractionError) as exc:
        extract(forged)
    assert exc.value.code == "DIMENSION_EXCEEDS_LIMITS"


def test_overstated_dimension_within_caps_yields_no_phantoms() -> None:
    data = make_xlsx({"Sales": [["A", "B"], ["1", "2"]]})
    forged = _patch_dimension(data, "A1:D10")
    result = extract(forged)
    assert [r.row_number for r in result.sheet.rows] == [2]
    assert result.stats.data_rows == 1


def test_missing_dimension_falls_back_to_actual_bounds() -> None:
    data = make_xlsx({"Sales": [["A", "B"], ["1", "2"], ["3", "4"]]})
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        names = zin.namelist()
        target = next(n for n in names if n.endswith("worksheets/sheet1.xml"))
        xml = zin.read(target).decode("utf-8")
        stripped, count = re.subn(_DIMENSION_TAG_RE, "", xml, count=1)
        assert count == 1
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, stripped if item.filename == target else zin.read(item.filename))
    result = extract(out.getvalue())
    assert [r.row_number for r in result.sheet.rows] == [2, 3]


def test_sparse_sheet_with_row_gaps() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["A", "B"])
    sheet.append(["1", "2"])
    sheet["A5"] = "x"
    sheet["B5"] = "y"
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = extract(buffer.getvalue())
    assert [r.row_number for r in result.sheet.rows] == [2, 5]
    assert result.sheet.rows[1].cells == ("x", "y")


def test_entity_declarations_refused() -> None:
    data = make_xlsx({"Sales": SALES_ROWS})
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        names = zin.namelist()
        target = next(n for n in names if n.endswith("worksheets/sheet1.xml"))
        xml = zin.read(target).decode("utf-8")
        assert "<worksheet" in xml
        poisoned = xml.replace("<worksheet", '<!DOCTYPE x [<!ENTITY x "y">]><worksheet', 1)
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, poisoned if item.filename == target else zin.read(item.filename))
    with pytest.raises(ExtractionError) as exc:
        extract(out.getvalue())
    assert exc.value.code == "FORBIDDEN_XML_ENTITIES"


# R3: sparse rows surface EmptyCells without `.row`; numbering must not crash.
def test_sparse_data_row_blank_first_cell() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet["A1"] = "A"
    sheet["B1"] = "B"
    sheet["B2"] = "val"  # A2 absent from XML: EmptyCell on read
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = extract(buffer.getvalue())
    assert result.sheet.headers == ("A", "B")
    assert len(result.sheet.rows) == 1
    assert result.sheet.rows[0].row_number == 2
    assert result.sheet.rows[0].cells == ("", "val")


def test_sparse_header_blank_first_cell_reports_typed_error() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet["B1"] = "B"  # A1 absent: EmptyCell header, not a crash
    sheet["B2"] = "val"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ExtractionError) as exc:
        extract(buffer.getvalue())
    assert exc.value.code == "BLANK_HEADER"


def test_sparse_multi_column_only_last_filled() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["A", "B", "C", "D"])
    sheet["D3"] = "z"  # entire row sparse except the last column
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = extract(buffer.getvalue())
    assert len(result.sheet.rows) == 1
    assert result.sheet.rows[0].row_number == 3
    assert result.sheet.rows[0].cells == ("", "", "", "z")
