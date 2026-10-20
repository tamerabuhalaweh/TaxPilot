"""REC-UI-01 demo report renderers (pure, no DB).

The PDF is produced by the dependency-free writer and the XLSX by openpyxl.
Both must carry the mandatory synthetic label and the deterministic findings /
evidence so an exported document can be verified against the run.
"""

from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.demo.fixtures import SYNTHETIC_LABEL, SYNTHETIC_LABEL_ASCII, get_fixture
from app.demo.pdf import Line, build_pdf
from app.demo.service import run_fixture
from app.demo.xlsx import build_xlsx

FIXTURE_ID = "synthetic-sales-gl-v1"


def test_pdf_is_structurally_valid_and_labelled() -> None:
    pdf = build_pdf(
        [
            Line(SYNTHETIC_LABEL_ASCII, size=12),
            Line("Synthetic fixture: synthetic-sales-gl-v1"),
        ]
    )
    assert pdf.startswith(b"%PDF-")
    assert pdf.rstrip().endswith(b"%%EOF")
    assert b"xref" in pdf
    assert SYNTHETIC_LABEL_ASCII.encode("ascii") in pdf
    assert b"synthetic-sales-gl-v1" in pdf


def test_pdf_paginates_without_losing_content() -> None:
    lines = [Line(f"line {index}") for index in range(200)]
    pdf = build_pdf(lines)
    assert pdf.startswith(b"%PDF-")
    assert pdf.count(b"/Type /Page ") >= 2
    assert b"line 199" in pdf


def test_xlsx_carries_label_findings_and_evidence() -> None:
    fixture = get_fixture(FIXTURE_ID)
    assert fixture is not None
    report, findings = run_fixture(fixture, tenant_id="t1", client_id="c1")
    data = build_xlsx(fixture, report, findings)
    assert data.startswith(b"PK")

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == ["Summary", "Findings", "Evidence"]
    for sheet in workbook.worksheets:
        assert sheet["A1"].value == SYNTHETIC_LABEL

    findings_sheet = workbook["Findings"]
    # Header row (3) + one row per finding.
    outcome_kinds = {
        findings_sheet.cell(row=row, column=5).value for row in range(4, 4 + len(findings.findings))
    }
    assert "DIFFERENCE" in outcome_kinds

    evidence_sheet = workbook["Evidence"]
    evidence_rows = [
        row for row in range(4, evidence_sheet.max_row + 1) if evidence_sheet.cell(row=row, column=1).value
    ]
    assert len(evidence_rows) == sum(len(item.evidence) for item in findings.findings)
