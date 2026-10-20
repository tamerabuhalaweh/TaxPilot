"""REC-UI-01 restricted synthetic-demo boundary (real PG + FastAPI).

Service-gating and harness mirror test_rec03_engagement_execution.py:
REQUIRE_SERVICES=1 (CI) turns unreachable services into FAILURE; all state lands
in the DEDICATED DISPOSABLE database (dbsafety-guarded). Synthetic data and
RFC-reserved example.com identities only.

Proves the demo boundary is opt-in, tenant-scoped, authenticated, and cannot be
mistaken for (or bypass) the real default-deny reconciliation path.
"""

import os
import sys
import uuid
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))  # repo root
sys.path.insert(0, os.path.dirname(__file__))  # sibling dbsafety module

import dbsafety  # noqa: E402

psycopg = pytest.importorskip("psycopg")

from app.demo.fixtures import SYNTHETIC_LABEL  # noqa: E402
from app.deps import settings_singleton  # noqa: E402
from app.main import create_app  # noqa: E402
from db.migrate import migrate  # noqa: E402

REQUIRE_SERVICES = os.environ.get("REQUIRE_SERVICES") == "1"
PASSWORD = "correct horse battery staple"
PERIOD = ("2026-01-01", "2026-12-31")
FIXTURE_ID = "synthetic-sales-gl-v1"


def _services_or_gate() -> None:
    try:
        with psycopg.connect(
            "postgresql://taxpilot:taxpilot@localhost:5432/taxpilot", connect_timeout=5
        ) as conn:
            conn.execute("SELECT 1")
        import redis as redis_lib

        redis_lib.Redis.from_url("redis://127.0.0.1:6379/0", socket_timeout=2).ping()
    except Exception as exc:
        if REQUIRE_SERVICES:
            pytest.fail(f"services unreachable with REQUIRE_SERVICES=1: {exc}")
        pytest.skip(f"services unreachable (documented optional local mode): {exc}")


_DURL: str | None = None


def _scratch_or_skip() -> str:
    _services_or_gate()
    global _DURL
    if _DURL is None:
        candidate = dbsafety.destructive_url()
        try:
            dbsafety.require_destructive_db(candidate)
        except ValueError as exc:
            if REQUIRE_SERVICES:
                pytest.fail(f"destructive-test safeguard refused target: {exc}")
            pytest.skip(f"destructive-test safeguard refused target: {exc}")
        migrate(candidate)
        _DURL = candidate
    else:
        migrate(_DURL)
    return _DURL


def _client(tmp_path, monkeypatch, *, enabled: bool):  # type: ignore[no-untyped-def]
    url = _scratch_or_skip()
    store_root = tmp_path / "quarantine"
    store_root.mkdir()
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.setenv("QUARANTINE_ROOT", str(store_root))
    monkeypatch.setenv("REDIS_URL", "redis://127.0.0.1:6379/0")
    if enabled:
        monkeypatch.setenv("TAXPILOT_DEMO_RECON_ENABLED", "1")
    else:
        monkeypatch.delenv("TAXPILOT_DEMO_RECON_ENABLED", raising=False)
        monkeypatch.delenv("DEMO_RECON_ENABLED", raising=False)
    settings_singleton.cache_clear()
    return TestClient(create_app(), raise_server_exceptions=False)


@pytest.fixture()
def api(tmp_path, monkeypatch):  # type: ignore[no-untyped-def]
    client = _client(tmp_path, monkeypatch, enabled=True)
    try:
        yield client
    finally:
        settings_singleton.cache_clear()


def _email(tag: str) -> str:
    return f"recui01-{tag}-{uuid.uuid4().hex[:8]}@example.com"


def _onboard(client: TestClient, tag: str) -> dict:
    email = _email(tag)
    reg = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": PASSWORD, "firm_name": "Synthetic Firm"},
    )
    assert reg.status_code == 201, reg.text
    token = reg.json()["verification_token"]
    assert client.post("/api/v1/auth/verify", json={"token": token}).status_code == 200
    logged = client.post(
        "/api/v1/auth/login", json={"email": email, "password": PASSWORD}
    )
    assert logged.status_code == 200, logged.text
    client.headers.update({"X-CSRF-Token": logged.json()["csrf_token"]})
    return {"email": email, "tenant_id": logged.json()["active_tenant_id"]}


def _engagement(client: TestClient) -> dict:
    created_client = client.post("/api/v1/clients", json={"name": "Synthetic Client"})
    assert created_client.status_code == 201, created_client.text
    engagement = client.post(
        "/api/v1/engagements",
        json={
            "client_id": created_client.json()["id"],
            "review_type": "SALES_GL",
            "period_start": PERIOD[0],
            "period_end": PERIOD[1],
        },
    )
    assert engagement.status_code == 201, engagement.text
    return engagement.json()


def _run(api: TestClient, engagement_id: str, fixture_id: str = FIXTURE_ID) -> dict:
    run = api.post(
        f"/api/v1/engagements/{engagement_id}/demo-recon-runs",
        json={"fixture_id": fixture_id},
    )
    assert run.status_code == 200, run.text
    return run.json()


def test_fixtures_and_run_happy_path(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)

    fixtures = api.get("/api/v1/demo-recon/fixtures")
    assert fixtures.status_code == 200, fixtures.text
    assert [f["fixture_id"] for f in fixtures.json()] == [
        FIXTURE_ID,
        "synthetic-sales-gl-v2",
    ]

    run = api.post(
        f"/api/v1/engagements/{engagement['id']}/demo-recon-runs",
        json={"fixture_id": FIXTURE_ID},
    )
    assert run.status_code == 200, run.text
    body = run.json()
    assert body["synthetic"] is True
    assert "SYNTHETIC DEMO DATA" in body["label"]
    difference = next(f for f in body["findings"] if f["outcome_kind"] == "DIFFERENCE")
    assert difference["delta"] == "100.000"
    assert difference["delta_currency"] == "JOD"
    assert body["inconclusive"] is True
    assert body["inconclusive_reasons"] == ["COVERAGE_DUPLICATE"]


def test_exports_are_labelled_binaries(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    base = f"/api/v1/engagements/{engagement['id']}/demo-recon-runs/report"

    pdf = api.get(base, params={"fixture_id": FIXTURE_ID, "format": "pdf"})
    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-")

    xlsx = api.get(base, params={"fixture_id": FIXTURE_ID, "format": "xlsx"})
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content.startswith(b"PK")

    bad = api.get(base, params={"fixture_id": FIXTURE_ID, "format": "docx"})
    assert bad.status_code == 422


def _engagement_with_period(api: TestClient, start: str, end: str) -> dict:
    created_client = api.post("/api/v1/clients", json={"name": "Period Client"})
    assert created_client.status_code == 201, created_client.text
    engagement = api.post(
        "/api/v1/engagements",
        json={
            "client_id": created_client.json()["id"],
            "review_type": "SALES_GL",
            "period_start": start,
            "period_end": end,
        },
    )
    assert engagement.status_code == 201, engagement.text
    return engagement.json()


def test_export_binds_to_displayed_run_and_states_periods(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement_with_period(api, "2026-02-01", "2026-06-30")
    run = _run(api, engagement["id"])
    # The calculation always uses the fixture period, never the engagement period.
    assert run["period_start"] == "2026-01-01"
    assert run["period_end"] == "2026-12-31"
    base = f"/api/v1/engagements/{engagement['id']}/demo-recon-runs/report"
    bound = {
        "fixture_id": FIXTURE_ID,
        "expected_recon_fingerprint": run["recon_fingerprint"],
        "expected_findings_fingerprint": run["findings_fingerprint"],
    }

    pdf = api.get(base, params={**bound, "format": "pdf"})
    assert pdf.status_code == 200, pdf.text
    body = pdf.content
    assert b"SYNTHETIC DEMO DATA - NOT CUSTOMER RECORDS / NOT A TAX FILING" in body
    assert b"Engagement period: 2026-02-01 to 2026-06-30" in body
    assert b"Fixture comparison period: 2026-01-01 to 2026-12-31" in body
    assert b"100.000" in body
    assert b"INV-1001" in body
    assert run["recon_fingerprint"].encode() in body
    assert run["findings_fingerprint"].encode() in body

    xlsx = api.get(base, params={**bound, "format": "xlsx"})
    assert xlsx.status_code == 200, xlsx.text
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Summary", "Findings", "Evidence"]
    for sheet in workbook.worksheets:
        assert sheet["A1"].value == SYNTHETIC_LABEL

    summary = workbook["Summary"]
    summary_rows = {
        summary.cell(row=row, column=1).value: summary.cell(row=row, column=2).value
        for row in range(3, summary.max_row + 1)
    }
    assert summary_rows["Engagement period"] == "2026-02-01 to 2026-06-30"
    assert summary_rows["Fixture comparison period"] == "2026-01-01 to 2026-12-31"
    assert summary_rows["Reconciliation fingerprint"] == run["recon_fingerprint"]
    assert summary_rows["Findings fingerprint"] == run["findings_fingerprint"]

    findings_sheet = workbook["Findings"]
    difference_rows = [
        row
        for row in range(4, findings_sheet.max_row + 1)
        if findings_sheet.cell(row=row, column=5).value == "DIFFERENCE"
    ]
    assert difference_rows, "XLSX must contain the JOD DIFFERENCE finding"
    difference_row = difference_rows[0]
    assert findings_sheet.cell(row=difference_row, column=13).value == "100.000"
    assert findings_sheet.cell(row=difference_row, column=14).value == "JOD"

    evidence_sheet = workbook["Evidence"]
    references = {
        evidence_sheet.cell(row=row, column=6).value
        for row in range(4, evidence_sheet.max_row + 1)
        if evidence_sheet.cell(row=row, column=6).value
    }
    assert "INV-1001" in references
    assert "J-1-1" in references


def test_export_refuses_a_different_run(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    run = _run(api, engagement["id"])
    base = f"/api/v1/engagements/{engagement['id']}/demo-recon-runs/report"

    wrong_findings = api.get(
        base,
        params={
            "fixture_id": FIXTURE_ID,
            "format": "pdf",
            "expected_findings_fingerprint": "not-the-displayed-run",
        },
    )
    assert wrong_findings.status_code == 409
    assert wrong_findings.json()["error"]["code"] == "STATE_CONFLICT"
    assert (
        wrong_findings.json()["error"]["details"]["reason"]
        == "DEMO_FINGERPRINT_MISMATCH"
    )

    wrong_recon = api.get(
        base,
        params={
            "fixture_id": FIXTURE_ID,
            "format": "xlsx",
            "expected_recon_fingerprint": "not-the-displayed-run",
            "expected_findings_fingerprint": run["findings_fingerprint"],
        },
    )
    assert wrong_recon.status_code == 409

    # A run of a different fixture can never authorise an export of this one.
    other = _run(api, engagement["id"], "synthetic-sales-gl-v2")
    crossed = api.get(
        base,
        params={
            "fixture_id": FIXTURE_ID,
            "format": "pdf",
            "expected_findings_fingerprint": other["findings_fingerprint"],
        },
    )
    assert crossed.status_code == 409
    assert crossed.json()["error"]["details"]["reason"] == "DEMO_FINGERPRINT_MISMATCH"


def test_fixture_switch_produces_distinct_results(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    first = _run(api, engagement["id"], FIXTURE_ID)
    second = _run(api, engagement["id"], "synthetic-sales-gl-v2")
    assert first["findings_fingerprint"] != second["findings_fingerprint"]
    assert len(first["findings"]) == 5
    assert len(second["findings"]) == 1
    difference = second["findings"][0]
    assert difference["outcome_kind"] == "DIFFERENCE"
    assert difference["delta"] == "50.000"
    assert difference["delta_currency"] == "JOD"


def test_unknown_fixture_is_rejected(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    run = api.post(
        f"/api/v1/engagements/{engagement['id']}/demo-recon-runs",
        json={"fixture_id": "nope"},
    )
    assert run.status_code == 422
    assert run.json()["error"]["code"] == "VALIDATION_ERROR"


def test_unauthenticated_is_denied(api: TestClient) -> None:
    anonymous = TestClient(api.app, raise_server_exceptions=False)
    assert anonymous.get("/api/v1/demo-recon/fixtures").status_code == 401
    denied = anonymous.post(
        f"/api/v1/engagements/{uuid.uuid4()}/demo-recon-runs",
        json={"fixture_id": FIXTURE_ID},
    )
    assert denied.status_code == 401


def test_cross_tenant_engagement_is_not_found(api: TestClient) -> None:
    _onboard(api, "owner-a")
    engagement = _engagement(api)

    intruder = TestClient(api.app, raise_server_exceptions=False)
    _onboard(intruder, "owner-b")
    run = intruder.post(
        f"/api/v1/engagements/{engagement['id']}/demo-recon-runs",
        json={"fixture_id": FIXTURE_ID},
    )
    assert run.status_code == 404
    assert run.json()["error"]["code"] == "NOT_FOUND"


def test_demo_is_disabled_by_default(tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    client = _client(tmp_path, monkeypatch, enabled=False)
    try:
        _onboard(client, "owner")
        engagement = _engagement(client)
        fixtures = client.get("/api/v1/demo-recon/fixtures")
        assert fixtures.status_code == 403
        assert fixtures.json()["error"]["code"] == "FORBIDDEN"
        assert fixtures.json()["error"]["details"]["reason"] == "DEMO_NOT_ENABLED"
        run = client.post(
            f"/api/v1/engagements/{engagement['id']}/demo-recon-runs",
            json={"fixture_id": FIXTURE_ID},
        )
        assert run.status_code == 403
    finally:
        settings_singleton.cache_clear()


def test_cookie_mutation_requires_csrf(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    url = f"/api/v1/engagements/{engagement['id']}/demo-recon-runs"
    missing = api.post(
        url, json={"fixture_id": FIXTURE_ID}, headers={"X-CSRF-Token": ""}
    )
    assert missing.status_code == 403
    assert missing.json()["error"]["code"] == "FORBIDDEN"
    invalid = api.post(
        url, json={"fixture_id": FIXTURE_ID}, headers={"X-CSRF-Token": "wrong"}
    )
    assert invalid.status_code == 403


def _add_viewer(owner: TestClient, tenant_id: str, tag: str) -> TestClient:
    """Create a second user, join the owner's tenant as a read-only viewer."""
    member = TestClient(owner.app, raise_server_exceptions=False)
    email = _email(tag)
    reg = member.post(
        "/api/v1/auth/register",
        json={"email": email, "password": PASSWORD, "firm_name": "Viewer Firm"},
    )
    assert reg.status_code == 201, reg.text
    assert (
        member.post(
            "/api/v1/auth/verify", json={"token": reg.json()["verification_token"]}
        ).status_code
        == 200
    )
    logged = member.post(
        "/api/v1/auth/login", json={"email": email, "password": PASSWORD}
    )
    assert logged.status_code == 200, logged.text
    member.headers.update({"X-CSRF-Token": logged.json()["csrf_token"]})
    invite = owner.post(
        "/api/v1/members/invitations", json={"email": email, "role": "viewer"}
    )
    assert invite.status_code == 201, invite.text
    accept = member.post(
        "/api/v1/members/invitations/accept", json={"token": invite.json()["token"]}
    )
    assert accept.status_code == 200, accept.text
    switched = member.post("/api/v1/auth/tenant", json={"tenant_id": tenant_id})
    assert switched.status_code == 200, switched.text
    return member


def test_viewer_cannot_run_demo(api: TestClient) -> None:
    owner = _onboard(api, "owner")
    engagement = _engagement(api)
    viewer = _add_viewer(api, owner["tenant_id"], "viewer")
    # Read-only roles may list fixtures but may not trigger a run.
    assert viewer.get("/api/v1/demo-recon/fixtures").status_code == 200
    denied = viewer.post(
        f"/api/v1/engagements/{engagement['id']}/demo-recon-runs",
        json={"fixture_id": FIXTURE_ID},
    )
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == "FORBIDDEN"


def test_real_recon_run_stays_default_denied(api: TestClient) -> None:
    _onboard(api, "owner")
    engagement = _engagement(api)
    activated = api.patch(
        f"/api/v1/engagements/{engagement['id']}", json={"status": "ACTIVE"}
    )
    assert activated.status_code == 200, activated.text
    blocked = api.post(
        f"/api/v1/engagements/{engagement['id']}/recon-runs",
        json={
            "sales_dataset_version_id": "syn-sales",
            "jofo_dataset_version_id": "syn-jofo",
            "idempotency_key": f"k-{uuid.uuid4().hex[:8]}",
        },
    )
    assert blocked.status_code == 409
    assert blocked.json()["blocked_reason"]
